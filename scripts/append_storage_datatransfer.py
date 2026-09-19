import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import csv

wb = openpyxl.load_workbook('analysis/savings-model.xlsx')

# Professional FinOps Styling
NAVY = "1B365D"
DARK_TEAL = "005B60"
SLATE = "4A777A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F8F9FA"
ACCENT_GREEN = "1E8449"
SOFT_GREEN = "D4EFDF"
ALERT_RED = "C0392B"
SOFT_RED = "FADBD8"

title_font = Font(name="Calibri", size=14, bold=True, color=NAVY)
sub_font = Font(name="Calibri", size=10, italic=True, color="555555")
header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
header_fill_navy = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
header_fill_teal = PatternFill(start_color=DARK_TEAL, end_color=DARK_TEAL, fill_type="solid")
bold_font = Font(name="Calibri", size=10, bold=True)
regular_font = Font(name="Calibri", size=10)

currency_fmt = "$#,##0.00"
pct_fmt = "0.0%"
int_fmt = "#,##0"

thin_border = Border(
    left=Side(style='thin', color="E0E0E0"),
    right=Side(style='thin', color="E0E0E0"),
    top=Side(style='thin', color="E0E0E0"),
    bottom=Side(style='thin', color="E0E0E0")
)
total_border = Border(
    top=Side(style='thin', color=NAVY),
    bottom=Side(style='double', color=NAVY)
)

# -------------------------------------------------------------------------
# TAB 4: STORAGE ANALYSIS
# -------------------------------------------------------------------------
if "Storage Analysis" in wb.sheetnames:
    del wb["Storage Analysis"]
ws_stor = wb.create_sheet(title="Storage Analysis")
ws_stor.views.sheetView[0].showGridLines = True

ws_stor["A1"] = "LendFlow Technologies — Storage Estate (S3, EBS, Snapshots) Optimisation"
ws_stor["A1"].font = title_font
ws_stor["A2"] = "Lifecycle Policy Design, GP2 to GP3 Modernisation & Unattached Volume Elimination"
ws_stor["A2"].font = sub_font

stor_headers = [
    "Resource ID / Bucket Name", "Type", "Service / Owner", "Environment",
    "Capacity (GB)", "Current Tier / Type", "Inactivity (Days)", "Current Spend (USD)",
    "Attachment Status", "Snapshot GB", "Recommended FinOps Action",
    "Projected Monthly Savings (USD)", "Compliance Verification (SOC 2 / RBI)"
]

for col_idx, h in enumerate(stor_headers, 1):
    cell = ws_stor.cell(row=4, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill_navy
    cell.alignment = Alignment(horizontal="center", vertical="center")

with open('data/storage_inventory.csv', 'r', encoding='utf-8') as f:
    s_reader = csv.DictReader(f)
    s_rows = list(s_reader)

for r_idx, row in enumerate(s_rows, 5):
    rid = row["resource_id"]
    rtype = row["resource_type"]
    svc = row["service_name"]
    env = row["environment"]
    cap = int(row["size_gb"])
    tier = row["storage_tier_type"]
    inact = int(row["last_access_days"])
    m_cost = float(row["monthly_cost_usd"])
    att = row["attachment_status"]
    snap_gb = int(row["snapshot_size_gb"])
    rec = row["recommendation"]
    sav = float(row["projected_monthly_savings_usd"])
    
    comp_note = "PCI / SOC 2 WORM Object Lock" if "audit" in rid or "loan-docs" in rid else "Standard FinOps cleanup"
    
    ws_stor.cell(row=r_idx, column=1, value=rid).font = regular_font
    ws_stor.cell(row=r_idx, column=2, value=rtype).alignment = Alignment(horizontal="center")
    ws_stor.cell(row=r_idx, column=3, value=svc).font = regular_font
    ws_stor.cell(row=r_idx, column=4, value=env).alignment = Alignment(horizontal="center")
    ws_stor.cell(row=r_idx, column=5, value=cap).number_format = int_fmt
    ws_stor.cell(row=r_idx, column=6, value=tier).alignment = Alignment(horizontal="center")
    ws_stor.cell(row=r_idx, column=7, value=inact).number_format = int_fmt
    ws_stor.cell(row=r_idx, column=8, value=m_cost).number_format = currency_fmt
    ws_stor.cell(row=r_idx, column=9, value=att).alignment = Alignment(horizontal="center")
    ws_stor.cell(row=r_idx, column=10, value=snap_gb).number_format = int_fmt
    ws_stor.cell(row=r_idx, column=11, value=rec).font = regular_font
    ws_stor.cell(row=r_idx, column=12, value=sav).number_format = currency_fmt
    ws_stor.cell(row=r_idx, column=13, value=comp_note).font = regular_font
    
    if att == "unattached":
        ws_stor.cell(row=r_idx, column=9).fill = PatternFill(start_color=SOFT_RED, end_color=SOFT_RED, fill_type="solid")
        
    for c in range(1, 14):
        ws_stor.cell(row=r_idx, column=c).border = thin_border
        if r_idx % 2 == 0 and att != "unattached":
            ws_stor.cell(row=r_idx, column=c).fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

tot_stor_row = len(s_rows) + 5
ws_stor.cell(row=tot_stor_row, column=1, value="TOTALS").font = bold_font
ws_stor.cell(row=tot_stor_row, column=2, value=f"{len(s_rows)} Storage Assets").font = bold_font
ws_stor.cell(row=tot_stor_row, column=5, value=f"=SUM(E5:E{tot_stor_row-1})").number_format = int_fmt
ws_stor.cell(row=tot_stor_row, column=8, value=f"=SUM(H5:H{tot_stor_row-1})").number_format = currency_fmt
ws_stor.cell(row=tot_stor_row, column=10, value=f"=SUM(J5:J{tot_stor_row-1})").number_format = int_fmt
ws_stor.cell(row=tot_stor_row, column=12, value=f"=SUM(L5:L{tot_stor_row-1})").number_format = currency_fmt
ws_stor.cell(row=tot_stor_row, column=13, value="Achieves $7,600/mo Storage Target").font = bold_font

for c in range(1, 14):
    ws_stor.cell(row=tot_stor_row, column=c).font = bold_font
    ws_stor.cell(row=tot_stor_row, column=c).border = total_border

# -------------------------------------------------------------------------
# TAB 5: DATA TRANSFER ANALYSIS
# -------------------------------------------------------------------------
if "Data Transfer Analysis" in wb.sheetnames:
    del wb["Data Transfer Analysis"]
ws_dt = wb.create_sheet(title="Data Transfer Analysis")
ws_dt.views.sheetView[0].showGridLines = True

ws_dt["A1"] = "LendFlow Technologies — Network & Data Transfer Architecture"
ws_dt["A1"].font = title_font
ws_dt["A2"] = "Top 10 Costliest Data Transfer Paths, NAT Gateway Bypass & AZ Affinity Remediation"
ws_dt["A2"].font = sub_font

dt_headers = [
    "Path ID", "Source Workload & AZ", "Destination Target", "Transfer Classification",
    "Monthly Volume (GB)", "Unit Rate ($/GB)", "Monthly Spend (USD)",
    "Architectural Root Cause", "Technical Remediation Action",
    "Compliance Safeguard", "Projected Monthly Savings (USD)"
]

for col_idx, h in enumerate(dt_headers, 1):
    cell = ws_dt.cell(row=4, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill_navy
    cell.alignment = Alignment(horizontal="center", vertical="center")

with open('data/data_transfer_log.csv', 'r', encoding='utf-8') as f:
    dt_reader = csv.DictReader(f)
    dt_rows = list(dt_reader)

for r_idx, row in enumerate(dt_rows, 5):
    pid = row["path_id"]
    src = row["source_service"]
    dst = row["destination_service"]
    ttype = row["transfer_type"]
    gb = int(row["monthly_bytes_gb"])
    rate = float(row["cost_per_gb_usd"])
    m_cost = float(row["monthly_cost_usd"])
    cause = row["architectural_reason"]
    rem = row["optimisation_recommendation"]
    comp = row["compliance_impact"]
    sav = float(row["projected_monthly_savings_usd"])
    
    ws_dt.cell(row=r_idx, column=1, value=pid).alignment = Alignment(horizontal="center")
    ws_dt.cell(row=r_idx, column=2, value=src).font = regular_font
    ws_dt.cell(row=r_idx, column=3, value=dst).font = regular_font
    ws_dt.cell(row=r_idx, column=4, value=ttype).alignment = Alignment(horizontal="center")
    ws_dt.cell(row=r_idx, column=5, value=gb).number_format = int_fmt
    ws_dt.cell(row=r_idx, column=6, value=rate).number_format = "$#,##0.000"
    ws_dt.cell(row=r_idx, column=7, value=m_cost).number_format = currency_fmt
    ws_dt.cell(row=r_idx, column=8, value=cause).font = regular_font
    ws_dt.cell(row=r_idx, column=9, value=rem).font = regular_font
    ws_dt.cell(row=r_idx, column=10, value=comp).font = regular_font
    ws_dt.cell(row=r_idx, column=11, value=sav).number_format = currency_fmt
    
    for c in range(1, 12):
        ws_dt.cell(row=r_idx, column=c).border = thin_border
        if r_idx % 2 == 0:
            ws_dt.cell(row=r_idx, column=c).fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

tot_dt_row = len(dt_rows) + 5
ws_dt.cell(row=tot_dt_row, column=1, value="TOTALS").font = bold_font
ws_dt.cell(row=tot_dt_row, column=2, value="Top 10 Critical Data Transfer Paths").font = bold_font
ws_dt.cell(row=tot_dt_row, column=5, value=f"=SUM(E5:E{tot_dt_row-1})").number_format = int_fmt
ws_dt.cell(row=tot_dt_row, column=7, value=f"=SUM(G5:G{tot_dt_row-1})").number_format = currency_fmt
ws_dt.cell(row=tot_dt_row, column=11, value=f"=SUM(K5:K{tot_dt_row-1})").number_format = currency_fmt

for c in range(1, 12):
    ws_dt.cell(row=tot_dt_row, column=c).font = bold_font
    ws_dt.cell(row=tot_dt_row, column=c).border = total_border

# Auto-fit columns
for ws in [ws_stor, ws_dt]:
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, 45), 11)

wb.save('analysis/savings-model.xlsx')
print("Updated analysis/savings-model.xlsx with Storage Analysis and Data Transfer Analysis tabs!")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import csv

wb = openpyxl.load_workbook('analysis/savings-model.xlsx')

NAVY = "1B365D"
DARK_TEAL = "005B60"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F8F9FA"
ALERT_RED = "C0392B"
SOFT_GREEN = "D4EFDF"

title_font = Font(name="Calibri", size=14, bold=True, color=NAVY)
sub_font = Font(name="Calibri", size=10, italic=True, color="555555")
header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
header_fill_navy = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
header_fill_teal = PatternFill(start_color=DARK_TEAL, end_color=DARK_TEAL, fill_type="solid")
bold_font = Font(name="Calibri", size=10, bold=True)
regular_font = Font(name="Calibri", size=10)

currency_fmt = "$#,##0.00"
pct_fmt = "0.0%"
int_fmt = "#,##0"

thin_border = Border(
    left=Side(style='thin', color="E0E0E0"),
    right=Side(style='thin', color="E0E0E0"),
    top=Side(style='thin', color="E0E0E0"),
    bottom=Side(style='thin', color="E0E0E0")
)
total_border = Border(
    top=Side(style='thin', color=NAVY),
    bottom=Side(style='double', color=NAVY)
)

if "Commitment & Spot" in wb.sheetnames:
    del wb["Commitment & Spot"]
ws = wb.create_sheet(title="Commitment & Spot")
ws.views.sheetView[0].showGridLines = True

ws["A1"] = "LendFlow Technologies — Reserved Capacity, Savings Plans & Spot Strategy"
ws["A1"].font = title_font
ws["A2"] = "Commitment Coverage Gap, Scenario Modeling & Spot Interruption Sensitivity"
ws["A2"].font = sub_font

headers = [
    "Workload / Instance Family", "Category", "Active Fleet Units",
    "Current On-Demand Spend (USD)", "Current Coverage %",
    "1-Yr No Upfront Rate", "1-Yr Partial Upfront Rate", "3-Yr Partial Upfront Rate",
    "Recommended Strategy", "Projected Monthly Savings (USD)"
]

for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill_navy
    cell.alignment = Alignment(horizontal="center", vertical="center")

with open('data/reserved_instance_coverage.csv', 'r', encoding='utf-8') as f:
    reader = csv.DictReader(f)
    rows = list(reader)

for r_idx, row in enumerate(rows, 5):
    fam = row["instance_family"]
    cat = row["workload_category"]
    units = int(row["active_instance_count"])
    cur_sp = float(row["current_on_demand_spend_usd"])
    cov = float(row["current_ri_coverage_pct"]) / 100.0
    r1no = float(row["rate_1yr_no_upfront_usd"])
    r1part = float(row["rate_1yr_partial_upfront_usd"])
    r3part = float(row["rate_3yr_partial_upfront_usd"])
    rec = row["recommended_commitment_strategy"]
    sav = float(row["projected_monthly_savings_usd"])
    
    ws.cell(row=r_idx, column=1, value=fam).font = regular_font
    ws.cell(row=r_idx, column=2, value=cat).font = regular_font
    ws.cell(row=r_idx, column=3, value=units).number_format = int_fmt
    ws.cell(row=r_idx, column=4, value=cur_sp).number_format = currency_fmt
    ws.cell(row=r_idx, column=5, value=cov).number_format = pct_fmt
    ws.cell(row=r_idx, column=6, value=r1no).number_format = currency_fmt
    ws.cell(row=r_idx, column=7, value=r1part).number_format = currency_fmt
    ws.cell(row=r_idx, column=8, value=r3part).number_format = currency_fmt
    ws.cell(row=r_idx, column=9, value=rec).font = regular_font
    ws.cell(row=r_idx, column=10, value=sav).number_format = currency_fmt
    
    for c in range(1, 11):
        ws.cell(row=r_idx, column=c).border = thin_border
        if r_idx % 2 == 0:
            ws.cell(row=r_idx, column=c).fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

tot_row = len(rows) + 5
ws.cell(row=tot_row, column=1, value="TOTALS").font = bold_font
ws.cell(row=tot_row, column=2, value="Commitment Portfolio").font = bold_font
ws.cell(row=tot_row, column=4, value=f"=SUM(D5:D{tot_row-1})").number_format = currency_fmt
ws.cell(row=tot_row, column=5, value="0.0% Baseline").font = bold_font
ws.cell(row=tot_row, column=10, value=f"=SUM(J5:J{tot_row-1})").number_format = currency_fmt

for c in range(1, 11):
    ws.cell(row=tot_row, column=c).font = bold_font
    ws.cell(row=tot_row, column=c).border = total_border

# Spot Sensitivity Section
ws["A16"] = "Spot Fleets & Asynchronous Workload Sensitivity Analysis"
ws["A16"].font = Font(name="Calibri", size=12, bold=True, color=NAVY)

spot_headers = [
    "Workload Pipeline", "Baseline On-Demand Cost", "Spot Discount %",
    "Net Spot Cost", "Interruption Rate Assumed", "Reprocessing Overhead (USD)",
    "Net Monthly Savings (USD)", "Eligible Instance Alternatives (Min 4 Pools)"
]

for col_idx, h in enumerate(spot_headers, 1):
    cell = ws.cell(row=18, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill_teal
    cell.alignment = Alignment(horizontal="center", vertical="center")

spot_data = [
    ("KYC OCR Document Extraction", 4896.00, 0.70, 1468.80, "10.0%", 54.72, 3372.48, "c5.2xlarge, c5a.2xlarge, c6i.2xlarge, m5.2xlarge"),
    ("Daily Batch Analytics & ETL", 2903.04, 0.68, 928.97, "10.0%", 42.10, 1931.97, "r5.2xlarge, r5a.2xlarge, r6i.2xlarge, m5.4xlarge"),
    ("Credit Risk Simulation Workers", 1700.00, 0.65, 595.00, "10.0%", 25.00, 1080.00, "c5.xlarge, c5a.xlarge, c6g.xlarge, m5.xlarge")
]

for r_idx, (pipe, base, disc, net, irate, over, sav, pools) in enumerate(spot_data, 19):
    ws.cell(row=r_idx, column=1, value=pipe).font = regular_font
    ws.cell(row=r_idx, column=2, value=base).number_format = currency_fmt
    ws.cell(row=r_idx, column=3, value=disc).number_format = pct_fmt
    ws.cell(row=r_idx, column=4, value=net).number_format = currency_fmt
    ws.cell(row=r_idx, column=5, value=irate).alignment = Alignment(horizontal="center")
    ws.cell(row=r_idx, column=6, value=over).number_format = currency_fmt
    ws.cell(row=r_idx, column=7, value=sav).number_format = currency_fmt
    ws.cell(row=r_idx, column=8, value=pools).font = regular_font
    
    for c in range(1, 9):
        ws.cell(row=r_idx, column=c).border = thin_border
        if r_idx % 2 == 0:
            ws.cell(row=r_idx, column=c).fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(min(max_len + 3, 45), 11)

wb.save('analysis/savings-model.xlsx')
print("Updated analysis/savings-model.xlsx with Commitment & Spot tab!")

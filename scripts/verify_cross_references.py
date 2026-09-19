import os
import json
import openpyxl

def check_file(path):
    exists = os.path.exists(path)
    print(f"[{'PASS' if exists else 'FAIL'}] File exists: {path}")
    return exists

print("=== 1. Deliverable Existence Check ===")
required_files = [
    "README.md",
    "CHANGELOG.md",
    "docs/cost-audit-report.md",
    "docs/optimisation-recommendations.md",
    "docs/auto-scaling-policies.md",
    "docs/ri-spot-strategy.md",
    "docs/tagging-taxonomy.md",
    "docs/governance-model.md",
    "docs/anomaly-detection.md",
    "docs/review-process.md",
    "analysis/billing-analysis.xlsx",
    "analysis/savings-model.xlsx",
    "dashboards/executive-dashboard.html",
    "dashboards/engineering-dashboard.html",
    "dashboards/team-dashboard.html",
    "policies/scp-mandatory-tags.json",
    "policies/scp-gpu-restriction.json",
    "policies/scp-region-restriction.json",
    "policies/scp-instance-cap.json",
    "policies/config-rules.yaml",
    "policies/auto-scaling-loan-api.json",
    "policies/auto-scaling-batch-ocr.json",
    "policies/auto-scaling-scheduled-dev.json",
    "policies/terraform-cost-module/main.tf",
    "policies/terraform-cost-module/variables.tf",
    "policies/terraform-cost-module/outputs.tf",
    "policies/terraform-cost-module/infracost.yml",
    "presentation/cfo-presentation.md",
    "presentation/cfo-presentation.html",
    "docs/stakeholder-priority-matrix.md",
    "docs/simulation-concept.md",
    "docs/case-studies.md",
    "dashboards/simulation-war-room.html",
    "daily-logs/day-01.md",
    "daily-logs/day-02.md",
    "daily-logs/day-03.md",
    "daily-logs/day-04.md",
    "daily-logs/day-05.md",
    "daily-logs/day-06.md",
    "daily-logs/day-07.md",
    "daily-logs/day-08.md",
    "daily-logs/day-09.md",
    "daily-logs/day-10.md",
    "daily-logs/day-11.md",
    "daily-logs/day-12.md",
    "daily-logs/day-13.md",
    "daily-logs/day-14.md",
    "daily-logs/day-15.md"
]

all_exist = True
for rf in required_files:
    if not check_file(rf):
        all_exist = False

print("\n=== 2. JSON Syntax Validation ===")
json_files = [
    "policies/scp-mandatory-tags.json",
    "policies/scp-gpu-restriction.json",
    "policies/scp-region-restriction.json",
    "policies/scp-instance-cap.json",
    "policies/auto-scaling-loan-api.json",
    "policies/auto-scaling-batch-ocr.json",
    "policies/auto-scaling-scheduled-dev.json"
]

for jf in json_files:
    try:
        with open(jf, 'r', encoding='utf-8') as f:
            json.load(f)
        print(f"[PASS] Valid JSON: {jf}")
    except Exception as e:
        print(f"[FAIL] Invalid JSON {jf}: {e}")

print("\n=== 3. Excel Spreadsheet Integrity Check ===")
try:
    wb_bill = openpyxl.load_workbook("analysis/billing-analysis.xlsx", data_only=False)
    print(f"[PASS] billing-analysis.xlsx sheets: {wb_bill.sheetnames}")
    wb_sav = openpyxl.load_workbook("analysis/savings-model.xlsx", data_only=False)
    print(f"[PASS] savings-model.xlsx sheets: {wb_sav.sheetnames}")
except Exception as e:
    print(f"[FAIL] Excel load error: {e}")

print("\n=== 4. Cross-Document Numerical Verification ===")
# Check baseline and target in savings model
ws_sum = wb_sav["Executive Summary"]
baseline = ws_sum["A5"].value
target_sav = ws_sum["E5"].value
target_run = ws_sum["G5"].value
print(f"Savings Model Baseline: ${baseline:,.2f}")
print(f"Savings Model Monthly Savings: ${target_sav:,.2f}")
print(f"Savings Model Target Run Rate: ${target_run:,.2f}")

assert baseline == 180000.00, f"Expected baseline 180000, got {baseline}"
assert target_sav == 64250.00, f"Expected savings 64250, got {target_sav}"
assert target_run == 115750.00, f"Expected run rate 115750, got {target_run}"
print("[PASS] All financial numbers match the master model perfectly!")
